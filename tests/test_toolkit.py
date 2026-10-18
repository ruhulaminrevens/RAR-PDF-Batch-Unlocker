"""Toolkit operation tests on synthetic files only."""
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import toolkit
from scripts.fixtures import make_pdf, KNOWN
from pypdf import PdfReader
from PIL import Image


def multipage(path, pages):
    make_pdf(path, None, pages=pages)
    return path


class ToolkitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.out = self.root / 'out'
        self.out.mkdir()

    def test_merge_counts_pages_and_refuses_single(self):
        a = multipage(self.root / 'a.pdf', 2)
        b = multipage(self.root / 'b.pdf', 3)
        res = toolkit.merge_pdfs([a, b], self.out / 'merged.pdf')
        self.assertEqual(res.pages, 5)
        self.assertEqual(len(PdfReader(str(res.outputs[0])).pages), 5)
        with self.assertRaises(toolkit.ToolkitError):
            toolkit.merge_pdfs([a], self.out / 'x.pdf')

    def test_merge_refuses_to_overwrite(self):
        a = multipage(self.root / 'a.pdf', 1)
        b = multipage(self.root / 'b.pdf', 1)
        target = self.out / 'm.pdf'
        toolkit.merge_pdfs([a, b], target)
        with self.assertRaisesRegex(toolkit.ToolkitError, 'already exists'):
            toolkit.merge_pdfs([a, b], target)

    def test_merge_requires_password_for_encrypted(self):
        enc = make_pdf(self.root / 'enc.pdf', KNOWN)
        plain = multipage(self.root / 'p.pdf', 1)
        with self.assertRaisesRegex(toolkit.ToolkitError, 'password-protected'):
            toolkit.merge_pdfs([enc, plain], self.out / 'm.pdf')
        res = toolkit.merge_pdfs([enc, plain], self.out / 'm.pdf', passwords={'enc.pdf': KNOWN})
        self.assertEqual(res.pages, 2)

    def test_split_per_page_and_ranges(self):
        src = multipage(self.root / 's.pdf', 4)
        res = toolkit.split_pdf(src, self.out)
        self.assertEqual(len(res.outputs), 4)
        self.assertTrue(all(p.exists() for p in res.outputs))
        res2 = toolkit.split_pdf(src, self.out / 'r', ranges=[(1, 2), (3, 4)])
        self.assertEqual(len(res2.outputs), 2)
        self.assertEqual(len(PdfReader(str(res2.outputs[0])).pages), 2)
        with self.assertRaises(toolkit.ToolkitError):
            toolkit.split_pdf(src, self.out / 'bad', ranges=[(1, 9)])

    def test_organize_reorder_delete_rotate(self):
        src = multipage(self.root / 'o.pdf', 5)
        res = toolkit.organize_pdf(src, self.out / 'o.pdf', keep=[5, 1, 2], rotations={5: 90})
        reader = PdfReader(str(res.outputs[0]))
        self.assertEqual(len(reader.pages), 3)
        self.assertEqual(reader.pages[0].rotation, 90)
        with self.assertRaises(toolkit.ToolkitError):
            toolkit.organize_pdf(src, self.out / 'x.pdf', keep=[])

    def test_parse_page_spec(self):
        self.assertEqual(toolkit.parse_page_spec('1,3-5,8', 10), [1, 3, 4, 5, 8])
        self.assertEqual(toolkit.parse_page_spec('2,2,2', 3), [2])
        for bad in ('0', '11', '3-2', 'a', '1-x'):
            with self.assertRaises(toolkit.ToolkitError):
                toolkit.parse_page_spec(bad, 10)

    def test_images_to_pdf_including_rgba_png(self):
        p1 = self.root / 'a.png'
        Image.new('RGB', (200, 120), (10, 80, 160)).save(p1)
        p2 = self.root / 'b.png'
        Image.new('RGBA', (120, 200), (200, 50, 50, 128)).save(p2)
        j = self.root / 'c.jpg'
        Image.new('RGB', (100, 100), (0, 0, 0)).save(j)
        res = toolkit.images_to_pdf([p1, p2, j], self.out / 'img.pdf')
        self.assertEqual(res.pages, 3)
        self.assertEqual(len(PdfReader(str(res.outputs[0])).pages), 3)
        with self.assertRaises(toolkit.ToolkitError):
            toolkit.images_to_pdf([self.root / 's.pdf'], self.out / 'bad.pdf')

    def test_pdf_to_images(self):
        src = multipage(self.root / 'r.pdf', 3)
        res = toolkit.pdf_to_images(src, self.out, fmt='png', dpi=100)
        self.assertEqual(len(res.outputs), 3)
        with Image.open(res.outputs[0]) as im:
            self.assertEqual(im.format, 'PNG')
            self.assertGreater(im.width, 100)
        res_j = toolkit.pdf_to_images(src, self.out / 'j', fmt='jpg', dpi=72, pages=[2])
        self.assertEqual(len(res_j.outputs), 1)
        self.assertTrue(res_j.outputs[0].name.endswith('p2.jpg'))
        for bad in (dict(dpi=10), dict(fmt='gif99')):
            with self.assertRaises(toolkit.ToolkitError):
                toolkit.pdf_to_images(src, self.out / 'b', **bad)

    def test_pdf_to_images_wrong_password(self):
        enc = make_pdf(self.root / 'e.pdf', KNOWN)
        with self.assertRaises(toolkit.ToolkitError):
            toolkit.pdf_to_images(enc, self.out, password='wrong')
        res = toolkit.pdf_to_images(enc, self.out / 'ok', password=KNOWN)
        self.assertEqual(len(res.outputs), 1)

    def test_pdf_to_docx_and_xlsx_handle_imageonly(self):
        # Blank synthetic pages have no selectable text; export must still succeed and say so.
        src = multipage(self.root / 'd.pdf', 2)
        res = toolkit.pdf_to_docx(src, self.out / 'd.docx')
        self.assertTrue(res.outputs[0].exists())
        self.assertIn('OCR', res.note)
        res_x = toolkit.pdf_to_xlsx(src, self.out / 'd.xlsx')
        self.assertTrue(res_x.outputs[0].exists())
        from openpyxl import load_workbook
        wb = load_workbook(res_x.outputs[0])
        self.assertIn('No tables found', wb.sheetnames)

    def test_formula_cell_guard(self):
        self.assertEqual(toolkit._safe_cell('=cmd'), "'=cmd")
        self.assertEqual(toolkit._safe_cell('ok'), 'ok')
        self.assertEqual(toolkit._safe_cell(None), '')


if __name__ == '__main__':
    unittest.main()
